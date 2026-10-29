"""Two renderers (LaTeX->PDF and python-docx->DOCX) for ONE list of content blocks.

Block types (tuples):
  ("title", title, subtitle, [meta lines])        cover page
  ("h1", text) ("h2", text) ("h3", text)
  ("p", text)                                    inline math as $...$ (LaTeX-compatible, simple symbols only)
  ("bullets", [text, ...])
  ("eq", latex)                                  display equation, auto-numbered; must be matplotlib-mathtext compatible
  ("table", caption, header, rows, align)        all cells are plain strings (inline $...$ allowed)
  ("fig", relative_path, caption, width_fraction)
  ("verb", text)                                 monospaced block
  ("pagebreak",)
"""
import os
import re
import subprocess
import hashlib

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

# ----------------------------------------------------------------------------- tex -> unicode (for DOCX)
GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π",
         "rho": "ρ", "sigma": "σ", "tau": "τ", "phi": "φ", "omega": "ω", "Delta": "Δ", "Sigma": "Σ", "Omega": "Ω", "Phi": "Φ", "psi": "ψ", "nu": "ν"}
SYMS = {"le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "times": "×", "approx": "≈", "infty": "∞", "cdot": "·", "pm": "±", "neq": "≠", "sim": "∼",
        "rightarrow": "→", "to": "→", "in": "∈", "sum": "Σ", "int": "∫", "partial": "∂", "nabla": "∇", "circledast": "⊛", "otimes": "⊗", "ll": "≪", "gg": "≫",
        "ldots": "…", "cdots": "⋯", "ast": "∗", "star": "⋆", "propto": "∝", "equiv": "≡", "degree": "°", "angle": "∠", "mid": "|", "quad": " ", "qquad": "  ",
        "log": "log", "ln": "ln", "exp": "exp", "sin": "sin", "cos": "cos", "min": "min", "max": "max", "arg": "arg", "sinc": "sinc", "Si": "Si", "mathrm": "",
        "bmod": " mod ", "mapsto": "↦", "ne": "≠", "gtrsim": "≳", "lesssim": "≲", "iff": "⇔", "lceil": "⌈", "rceil": "⌉", "kappa": "κ", "big": "", "Big": "", "circ": "°"}
SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
SUBL = {"i": "ᵢ", "j": "ⱼ", "k": "ₖ", "n": "ₙ", "m": "ₘ", "x": "ₓ", "r": "ᵣ", "s": "ₛ", "p": "ₚ", "L": "L", "H": "H"}


def tex_to_text(s):
    """Convert a simple LaTeX math fragment into Unicode text (used for DOCX inline math)."""
    s = re.sub(r"\\t?frac\{([^{}]*)\}\{([^{}]*)\}", lambda m: "(%s)/(%s)" % (m.group(1), m.group(2)) if len(m.group(1)) > 1 or len(m.group(2)) > 1 else "%s/%s" % (m.group(1), m.group(2)), s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\(?:mathrm|text|mathbf)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\hat\{([^{}]*)\}", lambda m: m.group(1) + "\u0302", s)
    s = re.sub(r"\\overline\{([^{}]*)\}", lambda m: m.group(1) + "\u0304", s)
    s = re.sub(r"\\left|\\right", "", s)
    s = re.sub(r"\\tfrac", r"\\frac", s)
    s = re.sub(r"\\bar\{([^{}]*)\}", lambda m: m.group(1) + "\u0304", s)
    s = re.sub(r"\\bar\s+([A-Za-z])", lambda m: m.group(1) + "\u0304", s)
    s = re.sub(r"\\mathcal\{?([A-Za-z])\}?", lambda m: {"F": "ℱ", "L": "ℒ"}.get(m.group(1), m.group(1)), s)
    s = re.sub(r"\^\{?\\circ\}?", "°", s)
    s = re.sub(r"\\([A-Za-z]+)", lambda m: GREEK.get(m.group(1), SYMS.get(m.group(1), m.group(1))), s)
    s = s.replace("\\,", " ").replace("\\;", " ").replace("\\ ", " ").replace("\\%", "%").replace("\\{", "{").replace("\\}", "}")

    def sup(m):
        t = m.group(1) if m.group(1) is not None else m.group(2)
        if t in ("*", "\\ast"):
            return "*"
        out = t.translate(SUP)
        return out if all(c in "0123456789+-=()n" for c in t) else "^(" + t + ")"

    def sub(m):
        t = m.group(1) if m.group(1) is not None else m.group(2)
        if all(c in "0123456789+-=()" for c in t):
            return t.translate(SUB)
        if len(t) == 1 and t in SUBL:
            return SUBL[t]
        return "_" + t if len(t) == 1 else "_(" + t + ")"
    s = re.sub(r"\^\{([^{}]*)\}|\^(\w)", sup, s)
    s = re.sub(r"_\{([^{}]*)\}|_(\w)", sub, s)
    return s.replace("{", "").replace("}", "")



INLINE_TOK = re.compile(r"(\$[^$]+\$|\*\*|\*|`)")


def inline_tokens(text):
    """Split text into (kind, string, bold, italic, code) tokens: kind 'math' ($...$, never altered by emphasis markers)
    or 'text'.  **bold**, *italic* and `code` toggle states; unbalanced markers raise."""
    out = []
    bold = ital = code = False
    for part in INLINE_TOK.split(text):
        if part == "":
            continue
        if code and part in ("**", "*"):
            out.append(("text", part, bold, ital, code))            # asterisks inside `code` are literal (e.g. results/*.json)
        elif part == "**":
            bold = not bold
        elif part == "*":
            ital = not ital
        elif part == "`":
            code = not code
        elif len(part) > 1 and part.startswith("$") and part.endswith("$"):
            out.append(("math", part[1:-1], bold, ital, code))
        else:
            out.append(("text", part, bold, ital, code))
    if bold or ital or code:
        raise ValueError("unbalanced emphasis markers in: " + text[:80])
    return out


def split_math(text):
    """Split text into [(is_math, string)] on $...$."""
    parts = re.split(r"(\$[^$]+\$)", text)
    return [(p.startswith("$") and p.endswith("$") and len(p) > 1, p[1:-1] if p.startswith("$") and p.endswith("$") and len(p) > 1 else p) for p in parts if p != ""]




def report_image(path, cache_dir=os.path.join(os.path.dirname(os.path.abspath(__file__)), "_figcache"), max_w=1900, quality=92):
    """Report-only copy of a figure: downscaled to <= max_w px wide and stored as high-quality JPEG (4:4:4) in a cache
    directory.  The original PNG in figures/ is never modified."""
    from PIL import Image
    os.makedirs(cache_dir, exist_ok=True)
    out = os.path.join(cache_dir, os.path.splitext(os.path.basename(path))[0] + ".jpg")
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(path):
        with Image.open(path) as src:
            im = src.convert("RGB")
        if im.width > max_w:
            im = im.resize((max_w, int(round(im.height * max_w / im.width))), Image.LANCZOS)
        im.save(out, "JPEG", quality=quality, subsampling=0, optimize=True)
    return out


def number_blocks(blocks):
    """Per-section numbering of tables and figures: ids like '2.3' (section = the first integer of the preceding h1 title,
    0 before the first numbered section).  Leading manual 'Table x.y.' prefixes of captions are stripped.
    Returns (ids, caps): ids[i] = '2.3' for table/figure blocks (None otherwise); caps[i] = caption without prefix."""
    ids, caps = [], []
    sec, cnt = 0, {"table": 0, "fig": 0}
    for b in blocks:
        if b[0] == "h1":
            m = re.search(r"\d+", b[1])
            sec = int(m.group(0)) if m else 0
            cnt = {"table": 0, "fig": 0}
        if b[0] in ("table", "fig"):
            cnt[b[0]] += 1
            ids.append("%d.%d" % (sec, cnt[b[0]]))
            cap = b[1] if b[0] == "table" else b[2]
            caps.append(re.sub(r"^Table \d+(\.\d+)?\.\s*", "", cap))
        else:
            ids.append(None)
            caps.append(None)
    return ids, caps


def check_references(blocks):
    """Return a list of problems: every 'Table x.y' / 'Figure x.y' mentioned in running text must exist."""
    ids, caps = number_blocks(blocks)
    have = {("Table" if b[0] == "table" else "Figure", i) for b, i in zip(blocks, ids) if i}
    bad = []
    for b in blocks:
        texts = []
        if b[0] in ("p", "h2", "h3"):
            texts = [b[1]]
        elif b[0] == "bullets":
            texts = b[1]
        for t in texts:
            for kind, num in re.findall(r"(Table|Figure)s? (\d+\.\d+)", t):
                if (kind, num) not in have:
                    bad.append("%s %s referenced but does not exist: ...%s" % (kind, num, t[:60]))
    return bad

# ----------------------------------------------------------------------------- LaTeX
TEX_ESC = {"&": r"\&", "%": r"\%", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}", "\\": r"\textbackslash{}",
           "<": r"\textless{}", ">": r"\textgreater{}", "|": r"\textbar{}"}
UNI_TEX = {"×": r"$\times$", "≈": r"$\approx$", "≤": r"$\le$", "≥": r"$\ge$", "±": r"$\pm$", "→": r"$\rightarrow$", "σ": r"$\sigma$", "μ": r"$\mu$", "α": r"$\alpha$",
           "β": r"$\beta$", "γ": r"$\gamma$", "θ": r"$\theta$", "λ": r"$\lambda$", "π": r"$\pi$", "φ": r"$\varphi$", "Δ": r"$\Delta$", "−": "-", "–": "--", "—": "---", "·": r"$\cdot$",
           "≠": r"$\neq$", "∞": r"$\infty$", "°": r"$^\circ$", "∼": r"$\sim$", "ρ": r"$\rho$", "τ": r"$\tau$", "ε": r"$\epsilon$", "ω": r"$\omega$", "≪": r"$\ll$", "∈": r"$\in$",
           "²": r"$^2$", "⁻": r"$^-$", "¹": r"$^1$", "√": r"$\sqrt{}$", "δ": r"$\delta$", "Σ": r"$\Sigma$", "∂": r"$\partial$", "∇": r"$\nabla$", "✓": r"\checkmark{}", "‰": r"\textperthousand{}"}


URL_RE = re.compile(r"https?://[A-Za-z0-9._~:/?\[\]@!$&'()*+,;=-]+")


def tex_text(text):
    out = []
    for kind, st, bold, ital, code in inline_tokens(text):
        if kind == "math":
            t = "$" + st + "$"
        else:
            t = "".join(TEX_ESC.get(c, UNI_TEX.get(c, c)) + (r"\allowbreak{}" if code and c in ",=/._(" else "") for c in st)
            if code:
                t = r"\texttt{" + t + "}"   # \allowbreak lets long identifiers wrap instead of running into the margin
                if URL_RE.fullmatch(st):
                    t = r"\href{%s}{%s}" % (st, t)   # clickable link in the PDF
            if bold:
                t = r"\textbf{" + t + "}"
            if ital:
                t = r"\textit{" + t + "}"
        out.append(t)
    return "".join(out)



def wrap_header(h, width=12):
    """Split a header at spaces (not inside $...$) into lines of about `width` characters, as a small centred tabular."""
    words, cur, inmath = [], "", False
    for ch in h:
        if ch == "$":
            inmath = not inmath
        if ch == " " and not inmath:
            words.append(cur)
            cur = ""
        else:
            cur += ch
    words.append(cur)
    lines, line = [], ""
    for w in words:
        if line and len(line) + 1 + len(w) > width:
            lines.append(line)
            line = w
        else:
            line = (line + " " + w) if line else w
    lines.append(line)
    if len(lines) == 1:
        return tex_text(h)
    return r"\begin{tabular}[b]{@{}c@{}}" + r"\\{}".join(tex_text(x) for x in lines) + r"\end{tabular}"


def to_latex(blocks, path_tex):
    L = [r"\documentclass[10pt,a4paper]{article}",
         r"\usepackage[margin=2.0cm]{geometry}\usepackage[T1]{fontenc}\usepackage{lmodern}\usepackage{amsmath,amssymb}",
         r"\usepackage{graphicx,booktabs,longtable,array,float,xcolor}",
         r"\graphicspath{{_figcache/}{../figures/}}\DeclareGraphicsExtensions{.jpg,.png}",   # report-size JPEGs when present, otherwise the original figures
         r"\usepackage[colorlinks=true,linkcolor=blue!50!black,pdftitle={Digital Image Processing, Assignment 2: Questions 1-6},pdfauthor={Shourya Sarkar}]{hyperref}",
         r"\setlength{\parskip}{3pt}\setlength{\parindent}{0pt}\makeatletter\long\def\@makecaption#1#2{\vskip 6pt\small\textbf{#1.} #2\par\vskip 4pt}\makeatother",
         r"\setcounter{secnumdepth}{0}", r"\begin{document}"]
    ids, caps = number_blocks(blocks)
    for bi, b in enumerate(blocks):
        t = b[0]
        if t == "title":
            _, title, sub, meta = b
            L += [r"\thispagestyle{empty}\begin{center}\vspace*{3cm}{\huge\bfseries %s\par}\vspace{0.6cm}{\Large %s\par}\vspace{1.5cm}" % (tex_text(title), tex_text(sub))]
            L += [r"{\large %s\par}\vspace{4pt}" % tex_text(m) for m in meta]
            L += [r"\end{center}\newpage"]
        elif t == "h1":
            L.append(r"\section{%s}" % tex_text(b[1]))
        elif t == "h2":
            L.append(r"\subsection{%s}" % tex_text(b[1]))
        elif t == "h3":
            L.append(r"\subsubsection{%s}" % tex_text(b[1]))
        elif t == "p":
            L.append(tex_text(b[1]) + "\n")
        elif t == "bullets":
            L.append(r"\begin{itemize}" + "".join(r"\item %s " % tex_text(x) for x in b[1]) + r"\end{itemize}")
        elif t == "eq":
            L.append(r"\begin{equation}%s\end{equation}" % b[1])
        elif t == "table":
            _, cap, header, rows, align = b
            ncol = len(header)
            al = align or ("l" + "r" * (ncol - 1))
            fs = r"\scriptsize" if ncol >= 8 else (r"\footnotesize" if ncol >= 5 else r"\small")
            L.append(r"\begin{table}[H]\centering%s{\small\textbf{Table %s.} %s\par}\vskip 3pt" % (fs, ids[bi], tex_text(caps[bi])))
            L.append(r"\resizebox{\ifdim\width>\linewidth\linewidth\else\width\fi}{!}{\begin{tabular}{%s}\toprule" % al)
            L.append(" & ".join(r"\textbf{%s}" % wrap_header(h) if len(header) >= 6 else r"\textbf{%s}" % tex_text(h) for h in header) + r"\\\midrule")
            for r in rows:
                L.append(" & ".join(tex_text(str(c)) for c in r) + r"\\")
            L.append(r"\bottomrule\end{tabular}}\end{table}")
        elif t == "fig":
            _, rel, cap, w = b
            L.append(r"\begin{figure}[H]\centering\includegraphics[width=%.2f\linewidth]{%s}\par\vskip 3pt{\small\textbf{Figure %s.} %s\par}\end{figure}" % (w, os.path.splitext(os.path.basename(report_image(os.path.join(ROOT, rel))))[0], ids[bi], tex_text(caps[bi])))
        elif t == "verb":
            L.append(r"{\footnotesize\begin{verbatim}" + "\n" + b[1] + "\n" + r"\end{verbatim}}")
        elif t == "pagebreak":
            L.append(r"\newpage")
    L.append(r"\end{document}")
    with open(path_tex, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


def build_pdf(blocks, outdir, name="DIP_A2_Report"):
    os.makedirs(outdir, exist_ok=True)
    tex = os.path.join(outdir, name + ".tex")
    to_latex(blocks, tex)
    logs = []
    for _ in range(2):
        r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", name + ".tex"], cwd=outdir, capture_output=True, text=True, encoding="utf-8", errors="replace")
        logs.append(r.returncode)
        if r.returncode != 0:
            print(r.stdout[-3000:])
            raise SystemExit("pdflatex failed")
    return os.path.join(outdir, name + ".pdf")


# ----------------------------------------------------------------------------- DOCX
def mathtext_safe(latex):
    """Map the LaTeX used in the report onto the subset understood by matplotlib's mathtext (for DOCX equation images
    only; the PDF keeps the original LaTeX)."""
    s = latex
    s = re.sub(r"\\xrightarrow\{[^{}]*\}", r"\\rightarrow", s)
    s = re.sub(r"\\text\{([^{}]*)\}", r"\\mathrm{\1}", s)
    s = re.sub(r"\\(?:Big|big|Bigg|bigg)[lr]?(?=[\(\)\[\]\|\{\}\\.])", "", s)
    s = re.sub(r"\\ge(?![a-zA-Z])", r"\\geq", s)
    s = re.sub(r"\\le(?![a-zA-Z])", r"\\leq", s)
    s = s.replace("\\circledast", "\\ast").replace("\\lesssim", "\\leq").replace("\\gtrsim", "\\geq")
    s = s.replace("\\;", "\\ ").replace("\\!", "")
    s = re.sub(r"\\frac([^{\s\\])", r"\\frac{\1}", s)                  # \frac12 -> \frac{1}{2} (mathtext needs braces)
    s = re.sub(r"(\\frac\{[^{}]*\})([^{\s\\])", r"\1{\2}", s)
    s = re.sub(r"\\mathcal\s+([A-Za-z])", r"\\mathcal{\1}", s)
    s = re.sub(r"\\mathbb\s*\{?([A-Za-z])\}?", r"\\mathbf{\1}", s)
    s = re.sub(r"\\(overline|hat|bar)\s+([A-Za-z])", r"\\\1{\2}", s)
    return s


def eq_png(latex, cache_dir):
    latex = mathtext_safe(latex)
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    os.makedirs(cache_dir, exist_ok=True)
    p = os.path.join(cache_dir, "eq_%s.png" % hashlib.md5(latex.encode()).hexdigest()[:12])
    if not os.path.exists(p):
        fig = plt.figure(figsize=(0.1, 0.1))
        fig.text(0, 0, "$" + latex + "$", fontsize=15)
        fig.savefig(p, dpi=200, bbox_inches="tight", pad_inches=0.04, transparent=False, facecolor="white")
        plt.close(fig)
    return p


def to_docx(blocks, path_docx, cache_dir):
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from PIL import Image

    doc = Document()
    cp = doc.core_properties
    cp.title = "Digital Image Processing, Assignment 2: Questions 1-6"
    cp.author = "Shourya Sarkar"
    cp.last_modified_by = "Shourya Sarkar"
    cp.subject = "Digital Image Processing, Assignment 2"
    cp.comments = ""
    cp.keywords = ""
    import datetime
    cp.created = cp.modified = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    text_w_cm = 17.0
    state = {"eq": 0, "tab": 0, "fig": 0}
    ids, caps = number_blocks(blocks)

    def add_runs(par, text, bold=False, size=None):
        for kind, st, b_, i_, c_ in inline_tokens(text):
            if kind == "math":
                r = par.add_run(tex_to_text(st))
                r.italic = True
            elif c_ and URL_RE.fullmatch(st):
                # clickable external hyperlink (blue, underlined, monospaced)
                r_id = par.part.relate_to(st, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), r_id)
                h.set(qn("w:history"), "1")
                wr = OxmlElement("w:r")
                rpr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Consolas")
                fonts.set(qn("w:hAnsi"), "Consolas")
                col = OxmlElement("w:color")
                col.set(qn("w:val"), "0563C1")
                und = OxmlElement("w:u")
                und.set(qn("w:val"), "single")
                for el in (fonts, col, und):
                    rpr.append(el)
                if size:
                    sz = OxmlElement("w:sz")
                    sz.set(qn("w:val"), str(int(size * 2)))
                    rpr.append(sz)
                wr.append(rpr)
                wt = OxmlElement("w:t")
                wt.text = st
                wr.append(wt)
                h.append(wr)
                par._p.append(h)
                continue
            else:
                r = par.add_run(st)
                if i_:
                    r.italic = True
                if c_:
                    r.font.name = "Consolas"
            r.bold = bool(bold or b_)
            if size:
                r.font.size = Pt(size)

    def caption(text, kind, number):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        state[kind] += 1
        r = p.add_run("%s %s. " % ("Table" if kind == "tab" else "Figure", number))
        r.bold = True
        r.font.size = Pt(9)
        add_runs(p, text, size=9)
        return p

    for bi, b in enumerate(blocks):
        t = b[0]
        if t == "title":
            _, title, sub, meta = b
            for _ in range(5):
                doc.add_paragraph()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(title); r.bold = True; r.font.size = Pt(26)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(sub); r.font.size = Pt(15)
            for _ in range(3):
                doc.add_paragraph()
            for m in meta:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, m, size=12)
            doc.add_page_break()
        elif t in ("h1", "h2", "h3"):
            h = doc.add_heading(level={"h1": 1, "h2": 2, "h3": 3}[t])
            add_runs(h, b[1])
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x2D, 0x5A)
        elif t == "p":
            p = doc.add_paragraph()
            add_runs(p, b[1])
        elif t == "bullets":
            for x in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, x)
        elif t == "eq":
            state["eq"] += 1
            png = eq_png(b[1], cache_dir)
            with Image.open(png) as im_eq:
                w, h = im_eq.size
            width_in = min(w / 200.0, 6.0)
            tbl = doc.add_table(rows=1, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c0, c1 = tbl.rows[0].cells
            c0.width, c1.width = Cm(15.0), Cm(2.0)
            c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c0.paragraphs[0].add_run().add_picture(png, width=Inches(width_in))
            c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            c1.paragraphs[0].add_run("(%d)" % state["eq"])
        elif t == "table":
            _, cap, header, rows, align = b
            caption(caps[bi], "tab", ids[bi])
            ncol = len(header)
            tbl = doc.add_table(rows=1, cols=ncol)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            fs = 7 if ncol >= 9 else (8 if ncol >= 6 else 9)
            for i, hdr in enumerate(header):
                c = tbl.rows[0].cells[i]
                c.paragraphs[0].text = ""
                add_runs(c.paragraphs[0], hdr, bold=True, size=fs)
                tcPr = c._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E8ECF4")
                tcPr.append(shd)
            for r_ in rows:
                cells = tbl.add_row().cells
                for i, v in enumerate(r_):
                    cells[i].paragraphs[0].text = ""
                    add_runs(cells[i].paragraphs[0], str(v), size=fs)
            for row in tbl.rows:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.paragraph_format.space_after = Pt(0)
            doc.add_paragraph()
        elif t == "fig":
            _, rel, cap, w = b
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(report_image(os.path.join(ROOT, rel)), width=Cm(text_w_cm * w))
            caption(caps[bi], "fig", ids[bi])
        elif t == "verb":
            p = doc.add_paragraph()
            r = p.add_run(b[1]); r.font.name = "Consolas"; r.font.size = Pt(8)
        elif t == "pagebreak":
            doc.add_page_break()
    doc.save(path_docx)
